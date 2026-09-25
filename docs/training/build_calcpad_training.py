"""Build the Calculation Pad training deck from the screenshots in docs/assets/screenshots.

    python -m pip install --target <temp folder> python-pptx
    set PYTHONPATH=<temp folder>
    python docs/training/build_calcpad_training.py

python-pptx is a documentation-only tool and deliberately not in the suite's locked
environment. Re-run this after re-capturing screenshots; the deck is regenerated whole.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
SHOTS = ROOT / "docs" / "assets" / "screenshots"
OUTPUT = Path(__file__).resolve().parent / "InnoCalc - Calculation Pad Training.pptx"
VERSION = "v0.0.1"

NAVY = RGBColor(0x12, 0x3B, 0x63)
INK = RGBColor(0x1C, 0x24, 0x31)
MUTED = RGBColor(0x5B, 0x66, 0x75)
SOFT = RGBColor(0xE6, 0xEE, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN, AMBER, RED = RGBColor(0x63, 0xBE, 0x7B), RGBColor(0xFF, 0xEB, 0x84), RGBColor(0xF8, 0x69, 0x6B)
WIDE, HIGH = Inches(13.333), Inches(7.5)

# Each slide: title, bullets, optional screenshot, speaker notes.
# A bullet written ("code", "...") is set in a monospace font.
SLIDES = [
    {"title": "What is the Calculation Pad?",
     "bullets": ["It is clever grid paper.",
                 "You write a sum. It does the maths.",
                 "It prints on the same Innovis sheet as every other calculation.",
                 "It remembers every version you save."],
     "image": "06-calculation-pad.png",
     "notes": "Think of the pad as squared paper that can do arithmetic. You write the "
              "working the way you would by hand, and the pad writes it out neatly with the "
              "numbers put in and the answer at the end."},
    {"title": "When do I use it?",
     "bullets": ["Load takedowns", "Design philosophy and assumptions", "Quick checks and assessments",
                 "Sketches and marked-up drawings",
                 "Not when there is already a module for it (Steel Member, Concrete Column)"],
     "notes": "If a proper module exists, use it: it has been tested and validated for that job. "
              "The pad is for everything else."},
    {"title": "Open the pad",
     "bullets": ["Open your project (Projects button).",
                 "Press New calculation.",
                 "Choose General on the left.",
                 "Click Calculation Pad."],
     "image": "05-choose-calculation-pad.png",
     "notes": "Press the star on the card to keep the Calculation Pad in Favourites."},
    {"title": "The three parts of the screen",
     "bullets": ["Left: every calculation in the project.",
                 "Middle: the boxes - who, what, and the list of cells.",
                 "Right: the printed sheet. It changes as you type."],
     "image": "06-calculation-pad.png",
     "notes": "Most of your time is spent on the right, typing into the sheet itself."},
    {"title": "Say what it is",
     "bullets": ["Calculation type and number - e.g. Lintel 0003",
                 "Package - which group of calculations it belongs to",
                 "Level - e.g. L02",
                 "Subject - one line that describes it",
                 "Description - words people will search for"],
     "image": "06-calculation-pad.png",
     "notes": "These fill the sheet header and file the calculation in the right folder. The "
              "Checked by box fills itself when a verifier signs the package off."},
    {"title": "Cells are like building blocks",
     "bullets": ["Section heading", "Narrative / assumptions (a line starting with - is a bullet)",
                 "Calculation", "Figure or sketch", "PDF drawing or Bluebeam markup",
                 "Blank A4 or A3 sheet for hand additions"],
     "image": "13-add-cells.png",
     "notes": "Add blocks with the + buttons in the middle panel. Use the arrows to move them up "
              "and down and x to delete one."},
    {"title": "Write a sum - in the sheet",
     "bullets": ["Click the grey box on the sheet.", "Type one sum per line:",
                 ("code", "w = 5.0 kN/m"), ("code", "L = 6.0 m"), ("code", "M_max = w * L^2 / 8   # kNm"),
                 "The answer appears underneath straight away."],
     "image": "07-type-in-the-sheet.png",
     "notes": "Each line becomes a printed line: the formula, the numbers put in, then the "
              "answer in bold. Later lines can use anything worked out above them."},
    {"title": "How to write it",
     "bullets": ["A few little rules:",
                 ("code", "L^2        power (or L**2)"),
                 ("code", "M_x_max    subscript: M with x,max below"),
                 ("code", "phi, alpha Greek letters by name"),
                 ("code", "# a note   a whole comment line prints as a note"),
                 ("code", "sqrt min max sin cos log pi ..."),
                 "Make up clear names - they are printed."],
     "notes": "The notation matches handcalcs, so it is the same as the Jupyter tools if you "
              "ever outgrow the pad."},
    {"title": "Numbers wear their units",
     "bullets": ["Write the unit straight after the number:",
                 ("code", "b = 250 mm"), ("code", "Z = b * d^2 / 6    # mm**3"),
                 "The # unit at the end converts the answer.",
                 "Add a force to a length and it goes red - that saves you from mistakes."],
     "image": "12-units-and-utilisation.png",
     "notes": "Units are real. 5 kN/m times 6 m squared divided by 8 really is 22.5 kNm. If the "
              "units cannot balance, the line says so instead of printing a wrong number. Set "
              "Units to Plain numbers if you only want labels."},
    {"title": "Declared values - your inputs",
     "bullets": ["A line that is just a number is an input: w = 5 kN/m",
                 "Inputs are listed in Declared variables in the middle panel.",
                 "Change one there and every sum below updates.",
                 "Nothing else moves on the sheet."],
     "image": "13-add-cells.png",
     "notes": "This is how you try a different span or load quickly without hunting through the "
              "sheet."},
    {"title": "Tell InnoCalc how hard it is working",
     "bullets": ["Name a line utilisation (or util):",
                 ("code", "M_Rd = 30 kNm"), ("code", "utilisation = M_max / M_Rd"),
                 "InnoCalc shows it in the Calculation Index and the package contents.",
                 "Over 1.0 is shown in red."],
     "image": "12-units-and-utilisation.png",
     "notes": "Without a utilisation line the pad is filed without one, which is fine for "
              "narrative-only sheets."},
    {"title": "Algebra when you need it",
     "bullets": ["Make a symbol, then let SymPy do the algebra:",
                 ("code", 'x = sym("x")'), ("code", "roots = solve(x^2 - 3*x - 10, x)"),
                 "Also: simplify, expand, factor, diff, integrate, limit",
                 "Tick Typeset equations for real fractions and roots on the sheet."],
     "notes": "SymPy does the algebra. The substitution line underneath still shows units."},
    {"title": "Drawings and markups",
     "bullets": ["Add a PDF drawing or Bluebeam markup cell.",
                 "Give the file and the pages (1, 1,3-5, or blank for all).",
                 "On screen it is a placeholder.",
                 "In a package the real pages are put in at that spot.",
                 "Insert blank A4 / A3 sheets for hand sketches."],
     "notes": "This is how a sketch or a marked-up drawing travels inside the calculation set, "
              "in the right place."},
    {"title": "Save",
     "bullets": ["Press SAVE (or Ctrl + S).",
                 "Yellow dot = changes not saved yet.",
                 "Each save is a new revision; the old one is kept as superseded.",
                 "The PDF is printed for you."],
     "image": "08-saved.png",
     "notes": "If you try to leave with unsaved changes InnoCalc asks whether to save them."},
    {"title": "Find it again",
     "bullets": ["Calculation Index lists everything in the project.",
                 "Search, or filter by package, level, module and type.",
                 "Tick Show superseded to see older revisions in grey.",
                 "Open brings it back to the pad."],
     "image": "02-calculation-index.png",
     "notes": "The PDF symbol opens the printed sheet; the folder symbol opens its folder."},
    {"title": "Too big for the pad?",
     "bullets": ["Export as Jupyter notebook - every sum becomes a handcalcs cell.",
                 "Check Pint and SymPy - see which libraries are installed.",
                 "Jupyter & handcalcs guide - links to training."],
     "image": "14-handcalcs-guide.png",
     "notes": "Use a notebook when a pad needs loops, tables or charts. If the same notebook is "
              "used again and again, it should become a proper module: see "
              "README-USER-CALCDEVELOPMENT.md."},
    {"title": "Why it will not run 'real' code",
     "bullets": ["Pads live on the shared drive and other people open them.",
                 "So a pad can only do arithmetic and the listed functions.",
                 "No imports, no files, no internet.",
                 "A pad can never harm someone else's computer."],
     "notes": "Every line is checked before it is worked out. Anything not allowed gets a "
              "message on the sheet instead."},
    {"title": "Golden rules",
     "bullets": ["One sum per line, clear names.", "Always write the units.",
                 "Add a utilisation line for design checks.",
                 "Write the assumptions in a narrative cell.",
                 "Save often - watch the yellow dot.",
                 "A pad is still checked like any hand calculation."],
     "notes": "The pad makes arithmetic reliable. It does not make the engineering right: the "
              "method, loads and assumptions are still yours to check."},
    {"title": "Try it: a simple beam",
     "bullets": ["New Calculation Pad, type Beam 0001.", "Write:",
                 ("code", "w = 5.0 kN/m"), ("code", "L = 6.0 m"),
                 ("code", "M_max = w * L^2 / 8   # kNm"), ("code", "M_Rd = 30 kNm"),
                 ("code", "utilisation = M_max / M_Rd"),
                 "Change L to 8 m in Declared variables. What happens? Save it."],
     "notes": "Answer: M_max goes from 22.5 to 40 kNm and utilisation from 0.75 to 1.33, shown "
              "in red. Put L back to 6 m, save, then find the calculation in the Calculation "
              "Index."},
    {"title": "Help and feedback",
     "bullets": ["Press Feedback in the top bar.", "Report a bug - say what you pressed and what went wrong.",
                 "Suggest an improvement - tell us what would help.",
                 "See something already listed? Press Me too.",
                 "User guide: README-ENDUSER.md"],
     "image": "09-feedback-page.png",
     "notes": "InnoCalc adds the page, module and version to your report automatically."},
]


def _text(frame, text, size, *, bold=False, colour=INK, font="Segoe UI", align=PP_ALIGN.LEFT):
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.name = Pt(size), bold, font
    run.font.color.rgb = colour


def _box(slide, left, top, width, height, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    item = slide.shapes.add_shape(shape, left, top, width, height)
    item.fill.solid()
    item.fill.fore_color.rgb = fill
    if line is None:
        item.line.fill.background()
    else:
        item.line.color.rgb = line
    item.shadow.inherit = False
    return item


def _brand(slide, left, top, size, *, inverse=False):
    mark = _box(slide, left, top, size, size, NAVY if inverse else WHITE,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    mark.adjustments[0] = 0.18
    frame = mark.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _text(frame, "IC", int(size / 12700 * 0.42), bold=True,
          colour=WHITE if inverse else NAVY, align=PP_ALIGN.CENTER)


def _fit(path: Path, max_width: int, max_height: int) -> tuple[int, int]:
    with Image.open(path) as image:
        ratio = image.width / image.height
    width = max_width
    height = int(width / ratio)
    if height > max_height:
        height = max_height
        width = int(height * ratio)
    return width, height


def _chrome(slide, title: str, number: int) -> None:
    _box(slide, 0, 0, WIDE, Inches(0.95), NAVY)
    _brand(slide, Inches(0.35), Inches(0.24), Inches(0.48))
    heading = slide.shapes.add_textbox(Inches(1.0), Inches(0.12), Inches(11.4), Inches(0.72))
    heading.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _text(heading.text_frame, title, 28, bold=True, colour=WHITE)
    footer = slide.shapes.add_textbox(Inches(0.35), Inches(7.05), Inches(12.6), Inches(0.35))
    _text(footer.text_frame, f"InnoCalc {VERSION}  |  Calculation Pad training  |  {number}", 10,
          colour=MUTED, align=PP_ALIGN.RIGHT)


def _bullets(slide, items, left, top, width, height) -> None:
    frame = slide.shapes.add_textbox(left, top, width, height).text_frame
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        code = isinstance(item, tuple)
        text = item[1] if code else item
        paragraph.space_after = Pt(4 if code else 10)
        run = paragraph.add_run()
        run.text = text if code else f"\u2022  {text}"
        run.font.name = "Consolas" if code else "Segoe UI"
        run.font.size = Pt(17 if code else 20)
        run.font.color.rgb = NAVY if code else INK
        if code:
            paragraph.level = 1


def title_slide(deck) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _box(slide, 0, 0, WIDE, HIGH, NAVY)
    _brand(slide, Inches(0.9), Inches(2.0), Inches(1.0))
    heading = slide.shapes.add_textbox(Inches(2.2), Inches(1.85), Inches(10), Inches(1.3))
    _text(heading.text_frame, "The Calculation Pad", 48, bold=True, colour=WHITE)
    sub = slide.shapes.add_textbox(Inches(2.2), Inches(3.05), Inches(10), Inches(0.8))
    _text(sub.text_frame, "Doing sums in InnoCalc - the easy way", 26, colour=SOFT)
    for offset, colour in enumerate((GREEN, AMBER, RED)):
        _box(slide, Inches(2.25 + offset * 0.55), Inches(4.2), Inches(0.4), Inches(0.12), colour)
    info = slide.shapes.add_textbox(Inches(2.2), Inches(4.6), Inches(10), Inches(0.6))
    _text(info.text_frame, f"InnoCalc {VERSION}  |  Innovis training", 16, colour=SOFT)
    slide.notes_slide.notes_text_frame.text = (
        "This session shows how to write a calculation in the Calculation Pad, from opening it "
        "to saving it and finding it again. No coding is needed. Bring a laptop with InnoCalc "
        "running if you want to follow along; the practice exercise is on the second-last slide.")


def content_slide(deck, spec: dict, number: int) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _chrome(slide, spec["title"], number)
    image = spec.get("image")
    if image and (SHOTS / image).is_file():
        _bullets(slide, spec["bullets"], Inches(0.5), Inches(1.3), Inches(4.9), Inches(5.6))
        width, height = _fit(SHOTS / image, Inches(7.4), Inches(5.6))
        left = Inches(5.6) + (Inches(7.4) - width) // 2
        top = Inches(1.3) + (Inches(5.6) - height) // 2
        frame = _box(slide, left - Emu(9525), top - Emu(9525), width + Emu(19050),
                     height + Emu(19050), WHITE, line=RGBColor(0xB8, 0xC2, 0xCF))
        frame.shadow.inherit = False
        slide.shapes.add_picture(str(SHOTS / image), left, top, width, height)
    else:
        _bullets(slide, spec["bullets"], Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.4))
    slide.notes_slide.notes_text_frame.text = spec["notes"]


def main() -> Path:
    deck = Presentation()
    deck.slide_width, deck.slide_height = WIDE, HIGH
    title_slide(deck)
    for number, spec in enumerate(SLIDES, start=2):
        content_slide(deck, spec, number)
    deck.core_properties.title = "InnoCalc - Calculation Pad training"
    deck.core_properties.author = "Innovis"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(main())
